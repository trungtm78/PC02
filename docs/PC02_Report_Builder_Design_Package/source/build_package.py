from pathlib import Path
import re,json,csv,html,zipfile
import yaml
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path('/workspace/scratch/c1f736b20146'); OUT=ROOT/'deliverables'; OUT.mkdir(exist_ok=True)
plan=(OUT/'APPROVED_PLAN.md').read_text()
parts={}
for chunk in re.split(r'(?m)(?=^## \d+\.)',plan):
 m=re.match(r'## (\d+)\.',chunk)
 if m: parts[int(m.group(1))]=chunk

# Một catalog duy nhất cho traceability, contract và tài liệu.
tasks=[
('M01-T01','Khảo sát source và chốt nghiệp vụ',[],['BRD','RBAC','repository'],3,4),
('M01-T02','Spike Excel và chọn engine',['M01-T01'],['parser','grid','formula','license'],4,6),
('M01-T03','Duyệt UX và mẫu nghiệm thu',['M01-T01'],['mockup','design system'],3,4),
('M02-T01','Schema phiên bản và migration',['M01-T02'],['database','storage'],3,4),
('M02-T02','Policy phân quyền báo cáo',['M02-T01'],['authorization','API'],2,3),
('M02-T03','Sinh kỳ và deadline',['M02-T01'],['scheduler','clock'],3,4),
('M03-T01','Parser token và kiểm tra workbook',['M01-T02','M02-T01'],['Excel import'],4,6),
('M03-T02','Wizard và xuất bản mẫu',['M03-T01','M02-T02','M02-T03','M01-T03'],['Setup UI','publish API'],4,5),
('M04-T01','Renderer lưới và editor theo kiểu',['M03-T01','M01-T03'],['Register UI','grid'],5,7),
('M04-T02','Lưu revision và hoàn thành',['M04-T01','M02-T02','M02-T03'],['submission API','autosave'],3,4),
('M04-T03','Mở lại và khóa lại',['M04-T02'],['grant API','dialog','audit'],2,3),
('M05-T01','Tổng hợp và công thức',['M04-T02','M03-T01'],['aggregate','formula engine'],4,6),
('M05-T02','Màn quản lý và nguồn số liệu',['M05-T01','M04-T03'],['manager UI','drilldown'],2,3),
('M06-T01','API và KPI tình trạng',['M04-T02','M04-T03'],['status query','indices'],3,4),
('M06-T02','Dashboard ma trận và lịch sử',['M06-T01','M01-T03'],['Status UI','history'],3,4),
('M06-T03','Xuất Excel theo snapshot',['M05-T01','M06-T01'],['export','download policy'],2,3),
('M07-T01','Tích hợp quyền thời gian và hồi quy',['M03-T02','M05-T02','M06-T02','M06-T03'],['IT','E2E','regression'],4,5),
('M07-T02','Hiệu năng và khả năng sử dụng',['M07-T01'],['benchmark','keyboard','responsive'],2,3),
('M08-T01','UAT và sửa lỗi nghiệm thu',['M07-T02'],['UAT','business signoff'],4,6),
('M08-T02','Pilot vận hành và rollback',['M08-T01'],['migration','monitoring','runbook'],2,3),
]
rows=[
('BR-01','Tải mẫu Excel','Upload .xlsx hợp lệ tạo bản nháp và preview; file sai loại/vượt giới hạn báo lỗi, không publish.','M03-T01','S02','file hợp lệ và file giả extension','parser trả schema hoặc lỗi có mã'),
('BR-01','Locked và input','Locked literal chỉ đọc; unlocked token tạo input; locked token và unlocked trống/literal báo Sheet!Cell.','M03-T01','S03','4 tổ hợp locked/token','schema field đúng và lỗi đúng ô'),
('BR-01','Grammar token','Chấp nhận TEXT NUM DATE TIME, format optional, AGG optional; Num chuẩn NUM, AGV cảnh báo alias AVG; token sai bị chặn.','M03-T01','S04','valid/invalid grammar và alias','fixture parser pass'),
('BR-01','Bố cục Excel','Giữ sheet merge font border wrap kích thước và freeze trong phạm vi hỗ trợ; cảnh báo thành phần ngoài phạm vi.','M04-T01','S11','golden workbook nhiều sheet và cột AK','đối soát vị trí và visual QA'),
('BR-01','Ô gộp và ẩn','Chỉ anchor của merge là field; input ẩn bị chặn publish; không nhân đôi giá trị.','M03-T01','S03','merge A2:C2 và hidden input','field count đúng'),
('BR-02','Lịch tuần','Chọn thứ và giờ, preview 6 kỳ; ISO week-year đúng tại tuần giao năm.','M02-T03','S05','tuần cuối năm/đầu năm','ngày giờ kỳ khớp fixture'),
('BR-02','Lịch tháng','Ngày 1–31/cuối tháng, ngày thiếu chuyển cuối tháng và preview rõ.','M02-T03','S06','tháng 2 nhuận và ngày 31','hạn đúng lịch'),
('BR-02','Lịch quý','Chọn tháng thứ 1/2/3 và ngày, quý hiện tại/kế tiếp; preview ngày thật.','M02-T03','S07','Q4 sang Q1','hạn không lệch năm'),
('BR-02','Ngày chính xác','Tạo một kỳ tại ngày giờ chọn, không tự lặp; hiển thị khoảng dữ liệu.','M02-T03','S08','ngày cụ thể với opens_at trước due_at','chỉ một kỳ được tạo'),
('BR-02','Khóa server','Tại now bằng hạn hoặc sau hạn API từ chối batch; đồng hồ client/job chậm không thay đổi kết quả.','M04-T02','S13','T−1ms T T+1ms','IT transaction và UI readonly'),
('BR-03','Phân công','Đúng một quản lý, nhiều người nhập không trùng, bỏ user inactive khi gán mới; publish thiếu người bị chặn.','M03-T02','S09','user trùng/inactive/không quản lý','validation + DB unique'),
('BR-03','Combo theo quyền','Chế độ nhập chỉ báo cáo được giao; quản lý chỉ báo cáo quản lý; truy cập trực tiếp ngoài phạm vi bị từ chối.','M02-T02','S11','ba identity và giả report id','API list/detail không lộ dữ liệu'),
('BR-04','Mặc định tổng hợp','Quản lý chọn báo cáo/kỳ mở tổng hợp all_saved, badge tạm tính và X/Y hoàn thành.','M05-T02','S15','3 người gồm nháp và hoàn thành','số tổng và contributor khớp'),
('BR-04','Xem cá nhân chỉ đọc','Quản lý chọn người thấy dữ liệu cùng kỳ, không có write; API sửa hộ bị chặn.','M05-T02','S16','quản lý thử PATCH dữ liệu người khác','403 và không đổi revision'),
('BR-05','Mở lại từng người','Grant chỉ đúng assignment/kỳ, lý do bắt buộc, manager hợp lệ; không mở các người hoặc kỳ khác.','M04-T03','S17','grant A nhưng thử ghi B','policy đúng phạm vi'),
('BR-05','Mặc định ba giờ','Default bằng server_now+3h, 23:30 chuyển 02:30 hôm sau; cấm thời điểm quá khứ.','M04-T03','S17','clock cố định 23:30','datetime dialog và API đúng'),
('BR-05','Grant hết hạn và thu hồi','API từ chối sau expires_at hoặc revoke nếu quá hạn gốc; trước hạn gốc vẫn áp dụng quyền gốc.','M04-T03','S14','expiry/revoke/before original due','quyền ghi đúng tại boundary'),
('BR-06','NUM','Editor và API từ chối chữ/NaN/infinity, giữ decimal, validate scale/min/max; không âm thầm làm tròn.','M04-T01','S12','âm thập phân overflow và format integer','raw decimal không sai lệch'),
('BR-06','DATE TIME','Có picker, từ chối 31/02 số bất kỳ 25:80; date-only không đổi ngày vì timezone.','M04-T01','S12','date leap year và time range','client/server cùng kết quả'),
('BR-06','TEXT và an toàn hiển thị','Giữ Unicode và số 0 đầu, escape HTML, TEXT bắt đầu = không thực thi công thức.','M04-T01','S11','text Unicode HTML và =1+1','UI/export đúng kiểu string'),
('BR-06','Paste','Paste TSV giữ vị trí; block có ô khóa/invalid từ chối nguyên khối với lỗi đúng ô.','M04-T01','S12','paste qua merge/locked/date sai','không có partial write'),
('BR-07','Autosave và mất mạng','Debounce 2s, trạng thái lưu rõ; retry idempotent; mất mạng không báo đã lưu.','M04-T02','S26','timeout trước/sau commit','một revision, không trùng'),
('BR-07','Hoàn thành','Chỉ hoàn thành khi required hợp lệ; sửa revision đã hoàn thành chuyển đang nhập; không khóa sớm ngoài hạn.','M04-T02','S11','complete thiếu/đủ và edit tiếp','state transition đúng'),
('BR-07','Đồng thời','Revision cũ trả 409; giữ buffer và cho so sánh, không overwrite người dùng âm thầm.','M04-T02','S26','hai tab cùng revision','một commit một conflict'),
('BR-04','SUM AVG và null','10,0,null cho SUM10 AVG5 COUNT2; empty SUM là null COUNT0; mỗi assignment góp một revision.','M05-T01','S18','fixture 3 assignment nhiều revision','expected independent pass'),
('BR-04','NONE và kiểu tổng hợp','TEXT DATE TIME mặc định NONE, bản tổng hiện không tổng hợp; cấu hình SUM text bị chặn.','M05-T01','S15','mixed types','không tự chọn giá trị người cuối'),
('BR-04','Công thức hai tầng','Tính cá nhân rồi tổng input rồi công thức tổng; tỷ lệ chung từ tử/mẫu tổng; cấm vòng lặp và external ref.','M05-T01','S18','tỷ lệ khác nhau + cycle + div0','kết quả đúng hoặc lỗi rõ'),
('BR-08','Theo dõi mọi lượt giao','C lấy assignment làm mẫu số, gồm người chưa nhập, loại miễn có lý do; người có 3 báo cáo là 3 lượt.','M06-T01','S19','3 reports cùng user và user inactive','mẫu số khớp phân công'),
('BR-08','Bộ lọc kỳ','Chọn report/kỳ/đơn vị/người/status, UI và export dùng cùng filter; mixed period hiện khoảng ngày thật.','M06-T02','S19','filter kết hợp và back từ B','filter và kết quả giữ đúng'),
('BR-08','Trạng thái và KPI','Progress, quyền nhập, đúng hạn độc lập; mở lại không xóa trễ; mẫu số0 hiện —.','M06-T01','S19','reopen completed và 0 obligations','KPI không cộng trùng'),
('BR-08','Drilldown và ma trận','Click status mở đúng B report/kỳ/người; không giao khác chưa nhập; ma trận giới hạn số cột.','M06-T02','S20','cặp có/không assignment','đúng route và trạng thái'),
('BR-09','Phiên bản mẫu','Publish mẫu mới không đổi kỳ cũ; không thay layout kỳ có dữ liệu; ngừng phát sinh không xóa dữ liệu.','M03-T02','S22','v1 đã nhập rồi publish v2','kỳ cũ vẫn render v1'),
('BR-09','Lịch sử và audit','Save/complete/grant/revoke/publish/phân công lưu actor thời điểm lý do và revision; không sửa audit từ UI.','M04-T03','S21','thực hiện chuỗi nghiệp vụ','audit đủ và đúng thứ tự'),
('BR-09','Xuất Excel','Export giữ layout/format/typed values và snapshot source; quyền download kiểm tra lại.','M06-T03','S25','export khi có concurrent save','file không trộn revision'),
('BR-10','Hai vai trò','User vừa quản lý vừa nhập có mode rõ; quản lý readonly, nhập chỉ own assignment.','M02-T02','S16','dual role','không có write hộ'),
('BR-10','Trạng thái UI','Mỗi màn có loading empty error forbidden; lock/mạng/conflict không che mất thông báo lưu.','M07-T01','S23','fault injection UI','ảnh và E2E state pass'),
('BR-10','Keyboard và responsive','Tab chỉ đi qua input, focus/error đọc được, 1366/1024/768 không che nút hoặc mất ô.','M07-T02','S24','keyboard và viewport','biên bản usability'),
('BR-10','Hiệu năng','Fixture 5000 input/200 người, 50 concurrent; p95 form3s save1s aggregate3s trên cấu hình đã ghi.','M07-T02','S11','benchmark sau chốt fixture','báo cáo p95 và không mất dữ liệu'),
('BR-10','Upload không an toàn','Reject macro/encrypted/external link/zip bomb, parser không gọi URL ngoài và có resource limit.','M03-T01','S04','fixture độc hại giả lập','lỗi an toàn và log không secret'),
('BR-02','Job và timezone','Sinh kỳ idempotent, phục hồi downtime không trùng, UTC storage hiển thị UTC+7 nhất quán.','M02-T03','S05','retry job và gap clock','unique period và hạn đúng'),
('BR-09','Thay phân công','Mặc định từ kỳ kế tiếp; kỳ hiện tại thao tác riêng có lý do và lịch sử mẫu số, không xóa submission.','M03-T02','S22','remove/add inactive/miễn','lịch sử giữ và số denominator đúng'),
('BR-10','Triển khai an toàn','Feature flag pilot, migration additive, restore/rollback đã thử; báo cáo cũ không hồi quy.','M08-T02','S23','staging deploy rồi rollback','runbook và biên bản restore'),
]
acs={};tests=[]
for n,(br,title,behavior,task,screen,scenario,evidence) in enumerate(rows,1):
 ac=f'AC-{n:03}'; fr=f'FR-{n:03}';test=f'UAT-{n:03}'; it=f'IT-{n:03}'
 acs[ac]={'requirement':title,'expected_behavior':behavior,'business_requirement':br,'functional_requirement':fr,'screen':screen,'tasks':[task], 'verification':{'integration':[it],'uat':[test]}}
 tests.append({'AC':ac,'FR':fr,'BR':br,'screen':screen,'task':task,'IT':it,'UAT':test,'scenario':scenario,'expected':behavior,'evidence':evidence})
contract={'contract_version':1,'status':'PROPOSED_BLOCKED_FOR_EXECUTION','source':{'specification':'BRD_PC02.docx + FRD_PC02.docx + APPROVED_PLAN.md','repository_baseline':'NOT PROVIDED'},'blockers':['Repository and deployment baseline','2–3 actual Excel templates','Business decisions D01–D10','Staging and separate role identities'],'acceptance_criteria':acs,'milestones':{},'gates':{'task':['acceptance_pass','targeted_tests_pass','affected_regression_pass','review_resolved','fresh_verification'],'milestone':['all_tasks_pass','full_suite_pass','alignment_pass','no_open_p0_p1'],'project':['all_milestones_pass','quality_checks_pass','uat_100_percent','independent_qa_pass','final_verification']}}
for tid,obj,deps,components,lo,hi in tasks:
 covered=[x for x in tests if x['task']==tid]
 testids=[x['IT'] for x in covered]+[x['UAT'] for x in covered]
 if not testids:testids=[{'M01-T01':'UAT-901','M01-T02':'IT-902','M01-T03':'UAT-903','M02-T01':'IT-904','M08-T01':'UAT-905'}.get(tid,'IT-906')]
 contract['milestones'].setdefault(tid[:3],{'tasks':{}})['tasks'][tid]={'objective':obj,'dependencies':deps,'components':components,'tests':testids,'estimate_person_days':[lo,hi],'red':'Viết fixture và expected độc lập cho '+obj.lower()+'.','green':'Triển khai vừa đủ cho hành vi và tiêu chí trong contract.','refactor':'Giảm trùng lặp, giữ nguyên test hành vi và chạy regression bị ảnh hưởng.','expected_evidence':'; '.join(x['evidence'] for x in covered) or 'Biên bản kiểm tra, artifact/version và môi trường xác minh.'}
(OUT/'EXECUTION_CONTRACT.yaml').write_text(yaml.safe_dump(contract,allow_unicode=True,sort_keys=False))
with (OUT/'TRACEABILITY.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(tests[0]));w.writeheader();w.writerows(tests)

business='''# BRD Báo cáo động PC02

Tài liệu xác định mục tiêu nghiệp vụ và phạm vi triển khai chức năng báo cáo theo mẫu Excel cho người thiết lập, người quản lý và người nhập. Kết quả cần đạt là người nhập làm việc trên biểu mẫu quen thuộc, quản lý xem số tổng có thể đối soát, và hệ thống cưỡng chế hạn nhập chính xác.

Phiên bản 1.0 ngày 09/10/2026. Chủ sở hữu nghiệp vụ và người phê duyệt: do đơn vị chỉ định. Trạng thái đề xuất chờ duyệt các quyết định D01 đến D10.

## Mục tiêu nghiệp vụ

| Mã | Mục tiêu và phạm vi |
|---|---|
| BR-01 | Tái sử dụng bố cục Excel để định nghĩa ô nhập mà không viết riêng từng màn báo cáo |
| BR-02 | Tự sinh kỳ và cưỡng chế hạn hoàn thành theo tuần tháng quý hoặc ngày cụ thể |
| BR-03 | Giao đúng người quản lý và nhiều người nhập theo từng báo cáo |
| BR-04 | Xem dữ liệu từng người và tổng hợp có nguồn gốc số liệu rõ |
| BR-05 | Cho phép nhập lại có thời hạn theo từng người và từng kỳ |
| BR-06 | Kiểm soát kiểu dữ liệu và định dạng ngay khi nhập |
| BR-07 | Lưu nháp an toàn và xác nhận hoàn thành tách biệt |
| BR-08 | Theo dõi đầy đủ nghĩa vụ báo cáo kể cả người chưa nhập |
| BR-09 | Giữ lịch sử mẫu dữ liệu phân công và xuất báo cáo đối soát |
| BR-10 | Tích hợp vào PC02 với quyền đúng tính ổn định và khả năng sử dụng phù hợp |

## Quy trình nghiệp vụ đích

Người thiết lập chuẩn bị file Excel và gán token cho ô unlocked. Sau khi kiểm tra preview, người thiết lập chọn lịch, phân công và xuất bản. Hệ thống sinh kỳ và các lượt giao. Người nhập chọn báo cáo và kỳ, lưu dữ liệu, kiểm tra và xác nhận hoàn thành. Khi tới hạn, server chặn mọi ghi không có grant còn hiệu lực. Người quản lý xem tổng hoặc từng người, có thể mở lại một lượt giao kèm lý do và ngày giờ khóa lại. Màn tình trạng luôn lấy toàn bộ lượt giao làm cơ sở để không bỏ sót người chưa nhập.

## Ranh giới phạm vi và giá trị

Phải có trong lần triển khai lõi: A B C, grammar token, lưu và tổng hợp, kỳ và khóa, mở lại, phân quyền, lịch sử và xuất Excel. Không yêu cầu thay thế các báo cáo nghiệp vụ đang sinh từ hồ sơ; không có luồng phê duyệt nhiều cấp, chữ ký số, lấy tự động số liệu hồ sơ, hay gửi email SMS trong phạm vi lõi. Các nội dung này chỉ thêm khi có yêu cầu thay đổi riêng.

Lợi ích dự kiến: giảm thời gian tạo màn nhập cho mẫu mới, giảm nhập sai kiểu, giảm tổng hợp thủ công và phát hiện chậm báo cáo sớm hơn. Chưa có baseline thực tế để cam kết phần trăm tiết kiệm. Pilot đo thời gian setup một mẫu, tỷ lệ lỗi input, thời gian đối soát và số lượt phải nhắc việc trước và sau.

## Các bên tham gia và trách nhiệm

| Vai trò | Trách nhiệm |
|---|---|
| Chủ nghiệp vụ | Chốt kỳ hạn ý nghĩa tổng hợp và nghiệm thu số liệu |
| Người thiết lập | Chuẩn bị mẫu kiểm tra schema lịch phân công và phiên bản |
| Người quản lý | Theo dõi xem dữ liệu và quyết định mở lại trong phạm vi |
| Người nhập | Nhập đúng dữ liệu tự kiểm tra và hoàn thành trước hạn |
| BA UX | Quản lý yêu cầu quyết định nghiệp vụ và tính dễ dùng |
| BE FE | Xây dựng theo FRD và thiết kế đã chốt giữ traceability |
| QA và người dùng UAT | Kiểm thử hành vi quyền thời gian số tổng và trải nghiệm |
| Vận hành | Triển khai sao lưu giám sát và phục hồi |

'''
business+=parts[1]+parts[2]+parts[3]+parts[11]+parts[14]+parts[15]

functional='''# FRD Báo cáo động PC02

Tài liệu mô tả chức năng quan sát được trên ba phân hệ và các quy tắc xử lý cần thống nhất giữa giao diện và máy chủ. Developer và QA sử dụng mã FR cùng AC và Screen ID để triển khai và nghiệm thu. Tất cả ví dụ người dùng và số liệu trong mockup là dữ liệu minh họa.

Phiên bản 1.0 ngày 09/10/2026. Áp dụng BRD cùng phiên bản. Quyết định đề xuất chưa duyệt được tập trung tại BRD mục D01 đến D10.

## Use case

| Mã | Actor và mục đích | Luồng và ngoại lệ |
|---|---|---|
| UC-01 | Thiết lập tạo và xuất bản báo cáo | S01→S02→S03→S05 đến S08→S09→S10; S04 nếu mẫu lỗi |
| UC-02 | Người nhập hoàn thành kỳ | S11→S12 nếu lỗi→S11 hoàn thành; S13 hết hạn; S26 mất mạng hoặc conflict |
| UC-03 | Quản lý xem tổng và cá nhân | S15→S18 nguồn số liệu hoặc S16 cá nhân; không cho sửa |
| UC-04 | Quản lý mở lại | S16→S17→S14 cho người nhập; hết hiệu lực trở về S13 |
| UC-05 | Theo dõi nghĩa vụ | S19→S20 ma trận→S16 chi tiết hoặc S21 lịch sử |
| UC-06 | Quản lý phiên bản | S01→S22→wizard nháp mới; giữ nguyên kỳ cũ |
| UC-07 | Xuất báo cáo | S15/S16/S19→S25; thất bại retry có kiểm tra quyền |

'''
functional+=parts[4]+parts[5]+parts[6]+parts[7]+parts[8]
functional+='\n## Danh mục chức năng và tiêu chí nghiệm thu\n\n'
for x in tests:
 functional+=f"### {x['FR']} {acs[x['AC']]['requirement']}\n\nNguồn {x['BR']}. Màn hình {x['screen']}. Nghiệm thu {x['AC']} qua {x['IT']} và {x['UAT']}.\n\n{x['expected']}\n\n"

basic='''# Basic Design Báo cáo động PC02

Tài liệu thiết kế tổng thể xác định các thành phần luồng dữ liệu và giao diện tích hợp cho module báo cáo động trong PC02. Thiết kế ưu tiên tái sử dụng nền tảng xác thực phân quyền lưu file và ghi nhật ký của hệ thống sau khi đối chiếu source code.

Phiên bản 1.0 ngày 09/10/2026. Kiến trúc logic độc lập framework; không coi tên endpoint hoặc entity đề xuất là cấu trúc đang tồn tại trong hệ thống.

## Các thành phần và ranh giới

| Thành phần | Trách nhiệm | Đầu ra |
|---|---|---|
| Report UI | Wizard grid manager dashboard state và keyboard | Request theo API contract và trạng thái người dùng |
| Report Application Service | Điều phối publish period assignment submission grant | Transaction và revision hợp lệ |
| Authorization Policy | Kiểm tra configurer manager owner scope | Allow hoặc lỗi quyền |
| Template Service | Lưu file immutable parse schema layout và validation | TemplateVersion đã kiểm tra |
| Clock Schedule Service | Sinh kỳ và tính effective edit permission | Period và server time |
| Calculation Service | Công thức giới hạn aggregate snapshot | Số tổng kèm tập nguồn |
| Status Query Service | Tổng hợp nghĩa vụ trạng thái và chỉ số | KPI bảng ma trận |
| Export Worker | Dựng file từ snapshot typed values | File và metadata kiểm soát quyền |
| Database và private file store | Dữ liệu phiên bản audit và file | Lịch sử có thể đối soát |

## Luồng dữ liệu chính

Upload đi vào Template Service trong vùng xử lý giới hạn tài nguyên, tạo schema đã kiểm tra. Xuất bản liên kết template version với lịch và phân công. Scheduler sinh kỳ cùng assignments. Grid lấy schema của kỳ và submission của người hiện tại. Save đi qua policy và clock, validate rồi tạo revision. Calculation Service đọc một tập revision nhất quán; Status Service đọc assignment và submission state. Export dùng đúng snapshot đã chọn, không đọc dữ liệu biến động nhiều lần trong cùng job.

## Quy tắc triển khai thành phần

Mặc định xây thành module trong backend hiện có để giữ transaction và quyền nhất quán. Không tách microservice chỉ vì tên logic khác nhau. Parser và export có thể chạy worker riêng trong cùng hạ tầng khi benchmark cho thấy cần. Formula engine chạy phía server làm nguồn kết quả; frontend preview không có quyền quyết định số chính thức. Nếu dùng cùng engine hai phía vẫn phải có test expected độc lập.

## Quan hệ dữ liệu

ReportDefinition có nhiều TemplateVersion và ScheduleVersion, có nhiều ReportPeriod. Mỗi Period tham chiếu chính xác một TemplateVersion và có nhiều Assignment. Một Assignment có tối đa một Submission, Submission có nhiều revision. Assignment có nhiều UnlockGrant lịch sử nhưng quyền hiện tại tính từ các grant chưa thu hồi còn hiệu lực. AggregateSnapshot tham chiếu Period và tập revision nguồn. ReportManagerAssignment lưu lịch sử quyền quản lý; quyền xem hiện tại dùng assignment quản lý hiện hành, lịch sử audit không tự cấp quyền truy cập.

'''
basic+=parts[3]+parts[8]+parts[9]+parts[10]+parts[12]+parts[14]

detail='''# Detail Design Báo cáo động PC02

Tài liệu hướng dẫn hiện thực các thuật toán dữ liệu API và chuyển trạng thái cho module báo cáo động. Dùng cùng FRD Basic Design mockup và execution contract phiên bản 1.0. Các tên route và kiểu dữ liệu bên dưới là hợp đồng đề xuất để ánh xạ sang framework hiện hành sau khảo sát repository.

## Quy ước định danh và phiên bản

ID public là chuỗi opaque do server sinh; client không suy ra ID từ tên hoặc vị trí. Field ID ổn định trong một TemplateVersion và liên kết sheet id cùng địa chỉ anchor. Không dùng A2 làm khóa toàn cục vì khác sheet và khác version có thể khác ý nghĩa. revision là số nguyên tăng đơn điệu trên một submission. Template schema version và calculation engine version lưu độc lập với số phiên bản báo cáo.

## Data dictionary và ràng buộc

| Entity | Field chính và kiểu logic | Constraint hoặc index |
|---|---|---|
| ReportDefinition | id string; code varchar64; name varchar255; status enum; effective_from date | unique code; index status |
| TemplateVersion | id; report_id FK; version int; file_ref; sha256; schema_json; parser_version; created_by; created_at timestamptz | unique report_id version; immutable sau publish |
| FieldDefinition | id; version_id FK; sheet_id string; address string; type enum; format string; agg enum; required bool; limits JSON | unique version sheet anchor; type và agg compatible |
| ReportPeriod | id; report_id FK; period_key varchar32; opens_at; original_due_at; period_start date; period_end date; template_version_id FK | unique report period_key; opens_at nhỏ hơn due_at |
| Assignment | id; period_id FK; input_user_id FK; unit_snapshot JSON; obligation enum; exemption_reason; assignment_version int | unique period user; index user period; giữ bản ghi khi miễn |
| Submission | id; assignment_id FK; current_revision bigint; progress enum; first_saved_at; first_completed_at; current_completed_at | unique assignment; index progress updated_at |
| SubmissionRevision | id; submission_id FK; revision bigint; values_json typed; complete bool; saved_by; committed_at | unique submission revision; immutable |
| UnlockGrant | id; assignment_id FK; start_at; expires_at; reason varchar1000; actor_id; revoked_at; revoked_by; revoke_reason | expires_at lớn hơn start_at; index assignment expires_at |
| AggregateSnapshot | id; period_id; mode enum; source_hash; source_revisions JSON; values JSON; engine_version; as_of | index period mode source_hash; không sửa snapshot |
| IdempotencyRecord | actor_id; action; key; request_hash; result_ref; status; expires_at | unique actor action key; TTL đề xuất24h |
| ExportJob | id; requested_by; scope_ref; snapshot_id; filters_json; status; file_ref; expires_at | quyền kiểm tra tại create và download; index status |
| AuditEvent | id; actor_id; action; entity_type; entity_id; before_revision; after_revision; reason; at; correlation_id | append-only; index entity at và actor at |

Các FK không cascade delete dữ liệu báo cáo đã phát sinh. Typed values có shape field_id→{type,value}; NUM value là chuỗi decimal chuẩn; null là giá trị rỗng thật, không phải chuỗi "null". Style schema không cho arbitrary HTML hay script. Giới hạn schema/file được kiểm tra trước lưu.

## Schema biểu mẫu mẫu

```json
{
  "schemaVersion": 1,
  "templateVersionId": "tv_demo_v1",
  "dateSystem": "1900",
  "locale": "vi-VN",
  "sheets": [{"id":"sheet_1","name":"BaoCao","merges":["B1:E1"]}],
  "fields": [
    {"id":"f_a2","sheetId":"sheet_1","address":"A2","type":"NUM","format":"#,##0","aggregate":"SUM","required":true,"limits":{"scale":0,"min":"0"}},
    {"id":"f_a5","sheetId":"sheet_1","address":"A5","type":"TEXT","aggregate":"NONE","required":false,"limits":{"maxLength":2000}},
    {"id":"f_ak12","sheetId":"sheet_1","address":"AK12","type":"DATE","format":"dd/mm/yyyy","aggregate":"NONE","required":true}
  ]
}
```

## Thuật toán import

1. Kiểm tra MIME extension kích thước và ZIP entries trước giải nén; bỏ đường dẫn traversal, giới hạn tổng kích thước và tỷ lệ giải nén.
2. Đọc workbook properties, date system, sheets, styles và effective protection. Chỉ parse dữ liệu không thực thi macro hoặc external connection.
3. Xác định vùng dùng hữu hạn và map merge anchor. Kiểm tra hidden input và feature không hỗ trợ.
4. Phân loại cell: literal locked, token unlocked, formula locked. Nếu bất nhất tạo lỗi cùng sheet/address/raw token đã escape.
5. Parse token theo grammar; normalize type/agg; validate format tương thích và rule limits.
6. Parse formula thành AST whitelist, xây dependency graph, kiểm tra vòng lặp và tham chiếu ngoài phạm vi. Không dùng eval.
7. Tạo schema/layout versioned, hash file và danh sách error/warning. Preview thể hiện lỗi; chỉ cho publish nếu error_count bằng0 và có ít nhất1 input.

## Parse token và số

Regex hình thức không thay validation ngữ nghĩa. Bóc đúng một cặp ngoặc ngoài, split tối đa3 phần theo |; type nằm trong allowlist. Không có agg thì NONE. Format rỗng dùng format chuẩn theo type. AGV chỉ alias khi parser option cho phép và phát warning. Escape không hỗ trợ trong grammar v1.

NUM client parse theo locale được cấu hình, không tự thử nhiều locale đến khi ra số. Backend chỉ nhận canonical decimal string có dấu trừ tùy chọn, phần nguyên và phần thập phân, không separator hoặc exponent. Kiểm tra tổng precision và scale trước chuyển decimal. Format hiển thị không sửa stored value. Comparison và aggregation dùng decimal, chỉ round tại lớp hiển thị hoặc hàm ROUND được định nghĩa.

## Quyền ghi và khóa giao dịch

```text
save(actor, assignment, expected_revision, key, patches):
  begin transaction
  load assignment + period + submission under appropriate lock
  require authenticated active actor owns this assignment
  check idempotency key hash; replay prior result with current permissions
  require obligation == REQUIRED and report permits existing-period writes
  require period.template_version == request.template_version
  now = database_clock_at_write_decision()
  allowed = opens_at <= now AND (
      now < original_due_at OR exists valid unrevoked grant at now)
  require allowed
  require current_revision == expected_revision
  validate every patch against field allowlist type limits and locked flag
  apply entire batch to current typed values
  recalculate supported formulas; fail batch on invalid computation policy
  if values actually changed: append revision with complete=false
  update submission pointer and audit atomically
  commit and return committed revision savedAt serverTime permissions
```

Idempotency replay phải vẫn xác thực actor và quyền đọc trước khi trả kết quả cũ. Cùng key khác payload trả IDEMPOTENCY_MISMATCH; không tái thực hiện. Write không thay đổi dữ liệu trả revision hiện tại, không hủy trạng thái hoàn thành. Complete cùng revision đã hoàn thành là no-op có phản hồi thành công. Check clock nằm trong giao dịch ngay trước ghi; nếu transaction phải chờ lock thì lấy lại now sau khi được lock.

## Request và response cụ thể

```json
PATCH /assignments/as_demo/values
{
  "templateVersionId":"tv_demo_v1",
  "expectedRevision":7,
  "idempotencyKey":"opaque-client-key",
  "patches":[{"fieldId":"f_a2","value":"125"},{"fieldId":"f_ak12","value":"2026-10-09"}]
}
```

```json
{
  "assignmentId":"as_demo",
  "revision":8,
  "progress":"IN_PROGRESS",
  "savedAt":"2026-10-09T07:10:00Z",
  "serverTime":"2026-10-09T07:10:00Z",
  "effectiveLockAt":"2026-10-09T10:00:00Z",
  "editable":true,
  "fieldErrors":[]
}
```

```json
HTTP 422
{
  "code":"CELL_VALIDATION",
  "message":"Có dữ liệu chưa hợp lệ",
  "errors":[{"fieldId":"f_ak12","sheet":"BaoCao","cell":"AK12","code":"DATE_INVALID","message":"Ngày không hợp lệ. Chọn ngày từ lịch."}],
  "correlationId":"request-reference"
}
```

POST unlock body gồm expiresAt theo ISO timezone offset và reason; server quyết định startsAt=now. Response gồm grantId originalDueAt effectiveLockAt và auditRef. POST complete body gồm expectedRevision và idempotencyKey. GET status query nhận reportId periodId hoặc period range cùng loại, unitId userId progress timing accessState, page pageSize sort; pageSize mặc định25 tối đa100, sort nằm trong allowlist. Response gồm items total summary asOf filterEcho. Thay đổi filter reset page về1.

## Chuyển trạng thái submission

| Trước | Sự kiện | Điều kiện | Sau |
|---|---|---|---|
| Chưa bắt đầu | Save có thay đổi | Được ghi và hợp lệ | Đang nhập |
| Đang nhập | Complete | Required đầy đủ không lỗi còn quyền ghi | Hoàn thành |
| Hoàn thành | Save không đổi | Request idempotent hoặc no-op | Hoàn thành |
| Hoàn thành | Save có đổi | Còn hạn hoặc grant hiệu lực | Đang nhập |
| Bất kỳ | Đến hạn | Không grant | Progress giữ nguyên quyền nhập Đã khóa |
| Bất kỳ | Grant | Quản lý đúng phạm vi | Progress giữ nguyên quyền Mở lại |
| Bất kỳ | Grant hết hạn | Đã quá hạn gốc | Progress giữ nguyên quyền Đã khóa |

## Thuật toán aggregate và status

Mở transaction đọc snapshot, lấy assignment REQUIRED trong period, chọn revision hiện tại theo mode all_saved hoặc completed. Lưu vector assignment→revision gồm cả chưa nhập=null. Với từng field NUM có agg, lọc null rồi tính decimal. NONE trả null kèm displayKind=not_aggregated. Formula tổng chạy trên các giá trị đã tổng hợp; null dependency trả thiếu dữ liệu thay vì0. Tạo source_hash từ period version mode và vector; cache theo hash. Khi save/grant/assignment thay đổi, aggregate hoặc status cache liên quan hết hiệu lực. Grant không đổi số tổng nhưng đổi metadata quyền/hạn.

Status query bắt đầu từ assignments LEFT JOIN submission để không mất người chưa nhập. Chọn tập kỳ theo filter; loại kỳ chưa mở khỏi KPI nghĩa vụ hiện tại và hiển thị riêng. completed_count dựa trên current revision. overdue_count dùng original_due_at và current_completed=false. reopened_count chỉ grant còn hiệu lực và đã quá hạn gốc. Query và export dùng cùng policy scope. Chụp asOf một lần để tất cả KPI và bảng cùng ranh giới thời gian.

## Frontend state và điều phối request

Khóa cache theo reportId periodId assignmentId templateVersionId. Đổi báo cáo/kỳ kiểm tra dirty buffer; xác nhận lưu hoặc bỏ; không trộn dữ liệu từ request cũ. Cancel request cũ hoặc kiểm tra request sequence trước apply response. Autosave tách dirty buffer pending batch và last persisted revision; serialize các batch trên cùng submission, không gửi song song nhiều expectedRevision bằng nhau. Response lỗi giữ dirty buffer; retry dùng cùng key chỉ khi payload không đổi. Complete chờ autosave flush thành công rồi dùng revision mới.

Hạn client dùng serverTime offset và refresh khi focus; display countdown không cấp quyền. Date picker trả date-only; browser timezone không đổi ngày. Validation phía client dùng cùng metadata schema nhưng server validate độc lập. Grid export/paste không dùng DOM làm nguồn dữ liệu; store typed values là nguồn UI.

## Export và bảo quản file

ExportJob chụp source revision vector trước khi enqueue; worker chỉ đọc revision bất biến. Xây workbook từ template immutable, thay token bằng giá trị typed và recalculated formulas/result theo policy; không để token trong file phát hành. Kèm sheet metadata hoặc thông tin đầu báo cáo về kỳ mẫu nguồn chế độ tổng hợp và thời điểm. TEXT xuất kiểu string, không diễn giải = + - @ như công thức. Lưu file private, download qua API kiểm tra quyền hiện hành, TTL đề xuất24h cho file export; template/submission retention theo chính sách đơn vị. Không dùng public permanent URL.

## Mã thông báo và kiểm thử lỗi

| Mã | Hành vi UI | Phục hồi |
|---|---|---|
| TEMPLATE_INVALID | Danh sách lỗi theo ô | Sửa Excel tải lại hoặc sửa thuộc tính nháp theo quyền |
| REPORT_LOCKED | Banner đỏ nhạt readonly và giữ phần chưa lưu | Liên hệ quản lý hoặc tải lại quyền sau mở lại |
| REVISION_CONFLICT | Dialog dữ liệu server và thay đổi chưa lưu | Tải bản mới rồi áp lại có kiểm tra |
| CELL_VALIDATION | Viền ô và drawer lỗi | Click lỗi focus đúng ô |
| FORBIDDEN | Trang không có quyền không tiết lộ số liệu | Quay lại danh sách được phép |
| PARSE_TIMEOUT | Trạng thái parse thất bại có mã tham chiếu | Tối giản mẫu hoặc điều chỉnh giới hạn có kiểm chứng |
| EXPORT_FAILED | Job failed không link file lỗi | Tạo lại từ snapshot còn hợp lệ |
| NETWORK_ERROR | Chưa lưu thử lại | Không đánh dấu completed hoặc saved |

## Kiểm thử đối soát bắt buộc

Fixture AGG01 có3 người: số lượng 10 0 null cho SUM10 AVG5 COUNT2. Fixture RATE01: người A xử lý9/10, người B xử lý1/90; tỷ lệ chung10/100=10%, không lấy trung bình90% và1,11%. Fixture DATE01 chứa workbook1900 và1904 cùng ngày hiển thị; ngày ISO phải giống nhau. Fixture LOCK01 chờ row lock vượt hạn trước write; request phải bị chặn. Fixture VERSION01 đổi mẫu v2 sau kỳ v1 có dữ liệu; render/export kỳ cũ phải giữv1.

Các test chưa chạy trên sản phẩm. Mỗi kết quả thực thi phải ghi commit môi trường clock fixture và bằng chứng; không lấy validation cấu trúc contract làm bằng chứng sản phẩm đã đạt.

'''
detail+=parts[13]+parts[15]

# WBS & test appendix giữ ngoài Word để đội phát triển sử dụng nhanh.
wbs='# Kế hoạch triển khai chi tiết\n\nƯớc lượng đề xuất '+str(sum(t[4] for t in tasks))+'–'+str(sum(t[5] for t in tasks))+' ngày công. Cập nhật sau M01.\n\n'
for tid,obj,deps,comps,lo,hi in tasks:
 t=contract['milestones'][tid[:3]]['tasks'][tid]
 wbs+=f'## {tid} {obj}\n\nPhụ thuộc: {", ".join(deps) or "Không"}. Thành phần: {", ".join(comps)}. Ước lượng {lo}–{hi} ngày công.\n\nRED: {t["red"]}\n\nGREEN: {t["green"]}\n\nREFACTOR: {t["refactor"]}\n\nTest: {", ".join(t["tests"])}. Bằng chứng: {t["expected_evidence"]}\n\n'
(OUT/'WBS.md').write_text(wbs)
uat='# Bộ tiêu chí và kịch bản nghiệm thu\n\nDùng tài khoản test và dữ liệu giả trên staging. Trạng thái tất cả: NOT RUN. Preconditions chung: phiên bản mẫu đã publish, kỳ và phân công xác định, clock server có thể điều khiển. Mỗi case tạo dữ liệu riêng; reset về baseline sau case.\n\n'
for x in tests:
 uat+=f'## {x["UAT"]} {acs[x["AC"]]["requirement"]}\n\nLiên kết {x["BR"]} → {x["FR"]} → {x["AC"]} → {x["task"]} → {x["screen"]}.\n\nThiết lập: {x["scenario"]}.\n\nBước1: mở màn {x["screen"]} với actor phù hợp chức năng. Bước2: thực hiện lần lượt các trường hợp đầu vào và điều kiện biên đã nêu. Bước3: đối chiếu UI, response API và revision/audit; tải lại trang để kiểm tra trạng thái lưu.\n\nExpected: {x["expected"]}\n\nEvidence: {x["evidence"]}. Kết quả: NOT RUN.\n\n'
uat+='''## Các gate kiểm tra bổ sung

UAT-901: biên bản D01 đến D10, repository baseline và role matrix được chốt; thiếu mục nào không cho qua M01.

IT-902: chạy parse render và tính công thức trên ba mẫu thật; đối soát toàn bộ ô input và license engine, ghi thời gian và giới hạn.

UAT-903: người thiết lập người nhập và quản lý walkthrough toàn bộ Screen ID; ký xác nhận luồng, không nhầm quyền và trạng thái.

IT-904: migration staging, foreign keys unique constraints revision concurrency và rollback tương thích; không mất dữ liệu cũ.

UAT-905: toàn bộ UAT-001 đến UAT-042 pass, không P0 P1, owner nghiệp vụ chấp thuận.
'''
(OUT/'UAT_CASES.md').write_text(uat)
(OUT/'PROGRESS.md').write_text('# PROGRESS\nUpdated: 2026-10-09 | Milestone: M00 | Task: 0/20\n\n## Completed\nKhảo sát giao diện chỉ đọc. BRD FRD Basic Design Detail Design và mockup đề xuất đã soạn.\n\n## In progress\nTask: Not started\nNEXT STEP: Chốt repository Excel thật và D01–D10.\n\n## Execution contract status\nTất cả AC: NOT RUN. Chưa bắt đầu implementation.\n\n## Test state\nTargeted: NOT RUN | Affected: NOT RUN | Full: NOT RUN\n\n## Current source state\nBranch: NOT PROVIDED | HEAD: NOT PROVIDED\n\n## Decisions\nCác mặc định trong BRD là đề xuất chưa được phê duyệt.\n')

def clean(s):return s.replace('**','').replace('`','')
def headingclean(s):return re.sub(r'[^\w\s]',' ',clean(s)).strip()
def setcell(cell,text,header=False):
 cell.text=clean(text);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 for p in cell.paragraphs:
  p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
  for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(255,255,255) if header else RGBColor(0,0,0);r.bold=header
def table(doc, lines):
 def splitrow(line):
  cells=[];buf='';depth=0
  for ch in line.strip().strip('|'):
   if ch=='{':depth+=1
   if ch=='}':depth=max(0,depth-1)
   if ch=='|' and depth==0:cells.append(buf.strip());buf=''
   else:buf+=ch
  cells.append(buf.strip());return cells
 data=[splitrow(l) for l in lines if not re.match(r'^\|[\s:|\-]+\|$',l)]
 cols=max(map(len,data));t=doc.add_table(rows=0,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 # content-aware proportions; short keys receive less width
 lens=[sum(min(len(row[i]) if i<len(row) else 0,180) for row in data)/len(data) for i in range(cols)]
 weights=[max(11,min(65,x)) for x in lens];total=sum(weights)
 widths=[6.9*x/total for x in weights]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for rownum,row in enumerate(data):
  cells=t.add_row().cells
  for i,c in enumerate(cells):
   c.width=Inches(widths[i]);setcell(c,row[i] if i<len(row) else '',rownum==0)
   tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173E65' if rownum==0 else ('F3F6F9' if rownum%2==0 else 'FFFFFF'));tcPr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for k in ['top','left','bottom','right']:
    e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   mar=OxmlElement('w:tcMar')
   for k in ['top','left','bottom','right']:
    e=OxmlElement('w:'+k);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tcPr.append(mar)
  if rownum==0:
   h=OxmlElement('w:tblHeader');t.rows[-1]._tr.get_or_add_trPr().append(h)
  # prevent split within individual row, allow whole table pagination
  ns=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(ns)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)

def make_doc(name,text):
 (OUT/(name+'.md')).write_text(text)
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
 for st in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  d.styles[st].font.name='Arial';d.styles[st].font.color.rgb=RGBColor(0,0,0)
 for style in d.styles:
  for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.line_spacing=1.12;d.styles['Normal'].paragraph_format.space_after=Pt(6)
 d.styles['Title'].font.size=Pt(24)
 for st,size in [('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
  d.styles[st].font.size=Pt(size);d.styles[st].paragraph_format.space_before=Pt(12);d.styles[st].paragraph_format.space_after=Pt(6)
 lines=text.splitlines();i=0;in_code=False;section_number=0
 while i<len(lines):
  l=lines[i]
  if l.startswith('```'):in_code=not in_code;i+=1;continue
  if in_code:
   p=d.add_paragraph(l);p.paragraph_format.space_after=Pt(0)
   for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8)
  elif l.startswith('|'):
   block=[]
   while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
   table(d,block);continue
  elif l.startswith('# '):d.add_paragraph(headingclean(l[2:]),'Title')
  elif l.startswith('### '):d.add_heading(headingclean(re.sub(r'^\d+(?:\.\d+)*[. ]*', '', l[4:])),2)
  elif l.startswith('## '):
   section_number+=1
   name_heading=re.sub(r'^\d+\.\s*','',l[3:])
   d.add_heading(str(section_number)+' '+headingclean(name_heading),1)
  elif l.strip():
   p=d.add_paragraph(clean(l))
   if l.startswith('Nguồn BR'):p.paragraph_format.keep_with_next=True
  i+=1
 foot=sec.footer.paragraphs[0];foot.alignment=2
 foot.add_run(name.replace('_',' ')+'  |  v1.0  |  ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 for r in foot.runs:r.font.size=Pt(8)
 for border in list(d.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.core_properties.title=headingclean(text.splitlines()[0][2:]);d.core_properties.author='';d.core_properties.subject='Báo cáo động PC02'
 d.save(OUT/(name+'.docx'))

for name,text in [('BRD_PC02',business),('FRD_PC02',functional),('Basic_Design_PC02',basic),('Detail_Design_PC02',detail)]:make_doc(name,text)
(OUT/'catalog.json').write_text(json.dumps({'tests':tests,'tasks':tasks},ensure_ascii=False,indent=2))
print(json.dumps({'documents':4,'AC':len(acs),'tasks':len(tasks),'estimate':[sum(t[4] for t in tasks),sum(t[5] for t in tasks)]}))
